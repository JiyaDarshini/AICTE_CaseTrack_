import io
import csv
from datetime import datetime, timezone
from typing import List, Dict, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
import logging

logger = logging.getLogger(__name__)


def generate_csv_report(data: List[Dict[str, Any]], fieldnames: List[str]) -> bytes:
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=fieldnames, extrasaction='ignore')
    writer.writeheader()
    for row in data:
        writer.writerow(row)
    return output.getvalue().encode('utf-8')


def generate_xlsx_report(title: str, headers: List[str], rows: List[List[Any]]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report"

    # Strict Palette & Font: Times New Roman, Orange headers (#F97316), light rows (#FFF7ED)
    title_font = Font(name="Times New Roman", size=16, bold=True, color="C2410C")
    header_font = Font(name="Times New Roman", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Times New Roman", size=10, color="1F2937")

    header_fill = PatternFill(start_color="F97316", end_color="F97316", fill_type="solid")
    alt_fill = PatternFill(start_color="FFF7ED", end_color="FFF7ED", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin', color='FED7AA'),
        right=Side(style='thin', color='FED7AA'),
        top=Side(style='thin', color='FED7AA'),
        bottom=Side(style='thin', color='FED7AA')
    )

    # Title row
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title_cell = ws.cell(row=1, column=1, value=f"AICTE CaseTrack — {title}")
    title_cell.font = title_font
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 35

    # Subtitle / Date
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    sub_cell = ws.cell(row=2, column=1, value=f"Generated on {datetime.now(timezone.utc).strftime('%d %B %Y, %H:%M UTC')}")
    sub_cell.font = Font(name="Times New Roman", size=9, italic=True, color="6B7280")
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    # Header row
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[3].height = 25

    # Data rows
    for row_idx, row_data in enumerate(rows, 4):
        is_even = (row_idx % 2 == 0)
        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=str(val) if val is not None else "")
            cell.font = data_font
            cell.border = thin_border
            if is_even:
                cell.fill = alt_fill
            cell.alignment = Alignment(vertical="center")

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_pdf_report(title: str, headers: List[str], rows: List[List[Any]]) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(letter),
        leftMargin=30,
        rightMargin=30,
        topMargin=30,
        bottomMargin=30
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'ReportTitle',
        parent=styles['Normal'],
        fontName='Times-Bold',
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#C2410C'),
        alignment=1,
        spaceAfter=6
    )
    meta_style = ParagraphStyle(
        'ReportMeta',
        parent=styles['Normal'],
        fontName='Times-Roman',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#6B7280'),
        alignment=1,
        spaceAfter=15
    )
    cell_style = ParagraphStyle(
        'Cell',
        parent=styles['Normal'],
        fontName='Times-Roman',
        fontSize=8,
        leading=10,
        textColor=colors.HexColor('#1F2937')
    )
    header_cell_style = ParagraphStyle(
        'HeaderCell',
        parent=styles['Normal'],
        fontName='Times-Bold',
        fontSize=9,
        leading=11,
        textColor=colors.white,
        alignment=1
    )

    story = [
        Paragraph(f"AICTE CaseTrack — {title}", title_style),
        Paragraph(f"Official Regulatory Report | Generated on {datetime.now(timezone.utc).strftime('%d %B %Y, %H:%M UTC')}", meta_style),
        Spacer(1, 10)
    ]

    # Convert table content to paragraphs for wrapping
    table_data = [[Paragraph(h, header_cell_style) for h in headers]]
    for r in rows:
        table_data.append([Paragraph(str(val) if val is not None else "", cell_style) for val in r])

    t = Table(table_data, repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#F97316')),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#FED7AA')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#FFF7ED')]),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))

    story.append(t)
    doc.build(story)
    return buf.getvalue()
